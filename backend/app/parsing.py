"""PDF/DOCX text extraction, section detection and basic field extraction."""
import re
from . import config

class ParseError(Exception):
    pass

HEAD = {
    "summary": ["summary", "professional summary", "profile", "objective", "career objective", "about me"],
    "skills": ["skills", "technical skills", "core competencies", "key skills", "skills & tools"],
    "experience": ["experience", "work experience", "professional experience", "employment history", "work history", "internship", "internships", "internship experience"],
    "education": ["education", "academic background", "academics"],
    "projects": ["projects", "academic projects", "personal projects", "key projects"],
    "certifications": ["certifications", "certificates", "licenses & certifications", "courses"],
}
LOOKUP = {h: k for k, v in HEAD.items() for h in v}
EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE = re.compile(r"(?<!\d)(\+?\d[\d\s().-]{8,}\d)(?!\d)")
DATE_RANGE = re.compile(r"((?:[A-Za-z]{3,9}\.?\s+)?(?:19|20)\d{2})\s*(?:-|–|—|to)\s*((?:[A-Za-z]{3,9}\.?\s+)?(?:19|20)\d{2}|present|current)", re.I)


def _clean(t):
    t = t.replace("\u2022", "- ").replace("\xa0", " ").replace("\r", "\n")
    return re.sub(r"[ \t]+", " ", t)


def _ocr(path):
    try:
        import pdfplumber, pytesseract
        with pdfplumber.open(path) as pdf:
            return "\n".join(pytesseract.image_to_string(p.to_image(resolution=200).original) for p in pdf.pages[: config.MAX_PAGES])
    except Exception:
        return None


def ocr_available():
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def read_file(path, ext):
    """Return (text, pages, ocr_used, tables)."""
    ocr_used, tables = False, 0
    if ext == "pdf":
        import pdfplumber
        try:
            with pdfplumber.open(path) as pdf:
                pages = len(pdf.pages)
                text = "\n".join((p.extract_text() or "") for p in pdf.pages[: config.MAX_PAGES])
        except Exception:
            raise ParseError("Could not read this PDF (it may be corrupt, encrypted or not a real PDF).")
        if len(text.strip()) < 40:
            ocr = _ocr(path)
            if ocr and len(ocr.strip()) >= 40:
                text, ocr_used = ocr, True
            else:
                raise ParseError("No extractable text found. The PDF may be a scan or image-only; OCR is not available or found nothing.")
    elif ext == "docx":
        import docx
        try:
            d = docx.Document(str(path))
        except Exception:
            raise ParseError("Could not read this DOCX (it may be corrupt or not a real Word document).")
        parts = [p.text for p in d.paragraphs]
        tables = len(d.tables)
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text, pages = "\n".join(parts), None
        if len(text.strip()) < 40:
            raise ParseError("The document contains almost no text.")
    else:
        raise ParseError("Unsupported file type.")
    return _clean(text), pages, ocr_used, tables


def split_sections(text):
    labelled, headings, unknown, cur, first = [], [], [], "header", True
    for raw in text.splitlines():
        l = raw.strip()
        if not l:
            continue
        key = re.sub(r"[:\-–|]+$", "", l.lower()).strip()
        if key in LOOKUP and len(l.split()) <= 4:
            cur = LOOKUP[key]; headings.append(l); first = False
            continue
        if (not first and l.isupper() and 1 <= len(l.split()) <= 3 and len(l) < 30 and not re.search(r"\d|@", l)):
            unknown.append(l); cur = "other"; continue
        first = False
        labelled.append((cur, l))
    sections = {}
    for s, l in labelled:
        sections.setdefault(s, []).append(l)
    return labelled, {k: "\n".join(v) for k, v in sections.items()}, headings, unknown


def extract(path, ext):
    text, pages, ocr_used, tables = read_file(path, ext)
    labelled, sections, headings, unknown = split_sections(text)
    email = (EMAIL.search(text) or [None])[0]
    phone = None
    for m in PHONE.finditer(text):
        if sum(c.isdigit() for c in m.group(1)) >= 10:
            phone = m.group(1).strip(); break
    name = None
    for s, l in labelled[:4]:
        if s == "header" and 1 <= len(l.split()) <= 5 and not re.search(r"\d|@|resume|curriculum", l, re.I) and sum(c.isalpha() for c in l) >= 3:
            name = l.title() if l.isupper() else l; break
    ranges = [m.group(0) for m in DATE_RANGE.finditer(text)]
    def lines(sec, n=10):
        return [re.sub(r"^[-*]\s*", "", x) for x in sections.get(sec, "").splitlines()][:n]
    skills_listed = [s.strip() for s in re.split(r"[,;|\n]", sections.get("skills", "")) if 1 < len(s.strip()) < 40][:40]
    return dict(
        text=text, pages=pages, ocr_used=ocr_used, tables=tables, labelled=labelled, sections=sections,
        headings=headings, unknown_headings=unknown, name=name, email=email, phone=phone, date_ranges=ranges,
        summary=sections.get("summary", "")[:500] or None, education=lines("education"),
        experience=lines("experience", 12), projects=lines("projects"), certifications=lines("certifications"),
        skills_listed=skills_listed)


def quality(p):
    """Descriptive parse-quality indicator (not a probability)."""
    n = len(p["text"].strip())
    pts = min(n / 1200, 1) * 40 + min(len([s for s in p["sections"] if s not in ("header", "other")]) / 4, 1) * 25
    pts += 10 * bool(p["email"]) + 5 * bool(p["phone"]) + 10 * bool(p["date_ranges"]) + 10 * bool(p["name"])
    warnings = []
    if p["ocr_used"]:
        warnings.append("OCR was used; recognition errors are possible.")
        pts = min(pts, 70)
    label = "Good text extraction" if pts >= 75 else "Partial text extraction" if pts >= 45 else "Limited text extraction"
    if p["ocr_used"]:
        label += " (OCR)"
    return {"score": round(pts), "label": label, "warnings": warnings}
